"""Import a PartsBox spreadsheet export.

The export has a "Komponenter" sheet (one row per part) and a "Förvaringplatser"
sheet (location id <-> Swedish name). Part rows reference a location by its
Swedish name, so we resolve name -> legacy_id -> our StorageLocation; that keeps
working after the locations were renamed to English.

Each part row becomes one Part plus one `add` StockEntry for the quantity at its
location. `meta` rows (PartsBox category placeholders) are skipped. Notes such as
"ca 75st ligger i SMD lådan" can't be split automatically — those rows are
returned in `review` so the user can split the stock by hand.

`run(db, path, commit)` returns a summary dict; with commit=False it rolls back.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Part, StockEntry, StorageLocation

SPLIT_HINT = re.compile(r"(ca|cirka|ungef)\.?\s*\d+\s*st|ligger\s+i\b|finns\s+ca", re.I)


def _num(v) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def _date(v) -> datetime | None:
    if isinstance(v, datetime):
        return v
    try:
        return datetime.fromisoformat(str(v)[:10])
    except (ValueError, TypeError):
        return None


def run(db: Session, path: str | Path, commit: bool = False) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    summary = {
        "created": 0,
        "updated": 0,
        "skipped_meta": 0,
        "stock_entries": 0,
        "locations_created": 0,
        "review": [],
        "warnings": [],
        "committed": commit,
    }

    # location name (Swedish, from export) -> our StorageLocation
    loc_by_legacy = {
        loc.legacy_id: loc
        for loc in db.scalars(select(StorageLocation)).all()
        if loc.legacy_id
    }
    sv_name_to_loc: dict[str, StorageLocation] = {}
    ws_loc = _sheet(wb, "förvaring")
    if ws_loc:
        for i, (lid, name) in enumerate(ws_loc.iter_rows(values_only=True)):
            if i == 0 or not name:
                continue
            loc = loc_by_legacy.get(str(lid).strip() if lid else None)
            if loc:
                sv_name_to_loc[str(name).strip()] = loc

    ws = _sheet(wb, "komponent")
    if ws is None:
        summary["warnings"].append("no 'Komponenter' sheet found")
        return summary

    header = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c).strip().lower() if c else "" for c in row]
            continue
        rec = dict(zip(header, row))
        name = _s(rec.get("namn"))
        if not name:
            continue
        typ = _s(rec.get("typ")).lower()
        if typ == "meta":
            summary["skipped_meta"] += 1
            continue

        mpn = _s(rec.get("mpn")) or name
        part = db.scalar(select(Part).where(Part.mpn == mpn)) if mpn else None
        notes = " | ".join(
            x for x in (_s(rec.get("kommentarer")), _s(rec.get("anteckningar"))) if x
        ) or None

        if part is None:
            part = Part(name=name, mpn=mpn)
            db.add(part)
            summary["created"] += 1
        else:
            summary["updated"] += 1
        part.description = _s(rec.get("beskrivning")) or part.description
        part.manufacturer = _s(rec.get("tillverkare")) or part.manufacturer
        part.footprint_raw = _s(rec.get("footprint")) or part.footprint_raw
        part.notes = notes or part.notes
        d = _date(rec.get("skapad"))
        if d:
            part.created_at = d
        db.flush()

        # stock
        qty = int(_num(rec.get("antal")) or 0)
        loc_name = _s(rec.get("förvaringsplas") or rec.get("förvaringsplats"))
        loc = sv_name_to_loc.get(loc_name)
        if loc_name and loc is None:
            loc = db.scalar(select(StorageLocation).where(StorageLocation.name == loc_name))
            if loc is None:
                loc = StorageLocation(name=loc_name)
                db.add(loc)
                db.flush()
                sv_name_to_loc[loc_name] = loc
                summary["locations_created"] += 1
        if qty:
            db.add(
                StockEntry(
                    part_id=part.id,
                    location_id=loc.id if loc else None,
                    delta=qty,
                    kind="add",
                    unit_price=_num(rec.get("pris (sek)")),
                    currency="SEK",
                    note="PartsBox import",
                )
            )
            summary["stock_entries"] += 1
        if notes and SPLIT_HINT.search(notes):
            summary["review"].append(
                {"part": name, "qty": qty, "location": loc_name, "note": notes}
            )

    if commit:
        db.commit()
    else:
        db.rollback()
    return summary


def _sheet(wb, prefix: str):
    for ws in wb.worksheets:
        if ws.title.strip().lower().startswith(prefix):
            return ws
    return None


def _s(v) -> str:
    return "" if v is None else str(v).strip()
